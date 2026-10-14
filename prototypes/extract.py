import openpyxl, pandas as pd, sys
src=sys.argv[1]
wv=openpyxl.load_workbook(src,data_only=True,read_only=True)
ws=wv['PENJUALAN']
rows=[]
for r in ws.iter_rows(min_row=4,max_col=53,values_only=True):
    if r[1] is None: continue
    rows.append(dict(unit=r[1],tipe=r[2],status=r[9],net=r[13],utj=r[45],batal=r[46],spr=r[47],cair=r[48],akad=r[49],ket=r[50]))
df=pd.DataFrame(rows)
for c in ['utj','batal','spr','cair','akad']: df[c]=pd.to_datetime(df[c],errors='coerce')
df.to_pickle('sales.pkl')
print(df.shape); print(df[['utj','batal','spr','cair','akad']].notna().sum())
print(df.status.value_counts())
# leads
ws=wv['LEADS']; L=[]
for r in ws.iter_rows(min_row=3,max_col=14,values_only=True):
    L.append((r[3],r[2],r[12]))
ld=pd.DataFrame(L,columns=['tgl','asal','status']); ld['tgl']=pd.to_datetime(ld.tgl,errors='coerce')
ld.to_pickle('leads.pkl'); print(ld.tgl.notna().sum(), ld.tgl.min(), ld.tgl.max())
